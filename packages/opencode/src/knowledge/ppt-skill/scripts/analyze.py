#!/usr/bin/env python3
"""读取 .pptx 的 slide layouts 与占位符目录（不解析正文），stdout 输出 JSON。"""
import json
import sys


def main() -> None:
    if len(sys.argv) < 2:
        print("usage: analyze.py <file.pptx>", file=sys.stderr)
        sys.exit(1)
    from pptx import Presentation

    prs = Presentation(sys.argv[1])
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        placeholders = []
        for ph in layout.placeholders:
            placeholders.append({
                "idx": ph.placeholder_format.idx,
                "type": str(ph.placeholder_format.type),
                "name": ph.name,
            })
        layouts.append({"index": i, "name": layout.name, "placeholders": placeholders})
    print(json.dumps({"layouts": layouts}, ensure_ascii=False))


if __name__ == "__main__":
    main()
